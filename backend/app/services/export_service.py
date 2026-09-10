from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import io

def generate_paper_excel(paper, extracted) -> io.BytesIO:
    """Generate an Excel file with one paper's summary."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Paper Summary"

    # Header row
    ws["A1"] = "Field"
    ws["B1"] = "Value"
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)

    rows = [
        ("Title", paper.title or "Not specified"),
        ("Filename", paper.filename),
        ("Research Problem", extracted.research_problem),
        ("Method Used", extracted.method_used),
        ("Dataset", extracted.dataset),
        ("Algorithm", extracted.algorithm),
        ("Results", extracted.results),
        ("Advantage", extracted.advantage),
        ("Limitation", extracted.limitation),
        ("Future Scope", extracted.future_scope),
        ("Narrative Summary", extracted.narrative_summary),
    ]

    for i, (field, value) in enumerate(rows, start=2):
        ws[f"A{i}"] = field
        ws[f"B{i}"] = value
        ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")

    # Widen columns for readability
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 80

    # Save to an in-memory buffer instead of disk
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_comparison_excel(papers_compared: list[dict], comparison_table: dict, similarity_scores: list[dict]) -> io.BytesIO:
    """Generate an Excel file with a full multi-paper comparison table."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    titles = [p["title"] for p in papers_compared]

    # Header row: Field, Paper 1, Paper 2, ...
    ws["A1"] = "Field"
    ws["A1"].font = Font(bold=True)
    for col, title in enumerate(titles, start=2):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)

    # One row per field
    field_labels = {
        "research_problem": "Research Problem",
        "method_used": "Method Used",
        "dataset": "Dataset",
        "algorithm": "Algorithm",
        "results": "Results",
        "advantage": "Advantage",
        "limitation": "Limitation",
        "future_scope": "Future Scope",
    }

    row_num = 2
    for field_key, field_label in field_labels.items():
        ws.cell(row=row_num, column=1, value=field_label)
        field_data = comparison_table.get(field_key, {})
        for col, title in enumerate(titles, start=2):
            value = field_data.get(title, "Not specified")
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_num += 1

    # Similarity scores on a second sheet
    ws2 = wb.create_sheet("Similarity Scores")
    ws2["A1"] = "Paper A"
    ws2["B1"] = "Paper B"
    ws2["C1"] = "Similarity Score"
    for cell in ["A1", "B1", "C1"]:
        ws2[cell].font = Font(bold=True)

    for i, s in enumerate(similarity_scores, start=2):
        ws2.cell(row=i, column=1, value=s["paper_a"])
        ws2.cell(row=i, column=2, value=s["paper_b"])
        ws2.cell(row=i, column=3, value=s["similarity_score"])

    # Column widths
    ws.column_dimensions["A"].width = 20
    for col_letter in "BCDEF"[:len(titles)]:
        ws.column_dimensions[col_letter].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer